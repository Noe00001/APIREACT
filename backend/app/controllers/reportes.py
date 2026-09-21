"""
reportes.py — Controlador REST de FastAPI para Generación de Reporte Diario de Ventas.
Exportación estructurada en PDF (ReportLab) y Excel XLSX (OpenPyXL).
Café Salomé / Café Cato.
"""
from io import BytesIO
from datetime import datetime, date
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.database import get_db
from app.models.models import Usuario, Venta, DetalleVenta
from app.auth import get_current_user, require_role

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])


def _obtener_ventas_dia(fecha_str: str, db: Session):
    try:
        fecha_obj = datetime.strptime(fecha_str, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato de fecha inválido. Debe ser YYYY-MM-DD")

    f_ini = datetime.combine(fecha_obj, datetime.min.time())
    f_fin = datetime.combine(fecha_obj, datetime.max.time())

    ventas = db.query(Venta).filter(
        Venta.fecha_hora >= f_ini,
        Venta.fecha_hora <= f_fin
    ).order_by(desc(Venta.fecha_hora)).all()

    items_desglosados = []
    total_dia = 0.0
    total_cant = 0

    for v in ventas:
        cliente_nom = f"{v.cliente.nombre} {v.cliente.apellido}" if v.cliente else "Cliente General"
        hora_str = v.fecha_hora.strftime("%H:%M") if v.fecha_hora else ""
        for det in v.detalles:
            items_desglosados.append({
                "fecha": fecha_str,
                "hora": hora_str,
                "numero_venta": v.numero_venta,
                "cliente": cliente_nom,
                "item": det.nombre_item,
                "tipo": det.tipo_item,
                "cantidad": det.cantidad,
                "precio_unitario": float(det.precio_unitario),
                "subtotal": float(det.subtotal),
                "total_venta": float(v.total),
                "metodo_pago": v.metodo_pago,
                "estado": v.estado,
            })
            total_cant += det.cantidad
        total_dia += float(v.total)

    return {
        "fecha": fecha_str,
        "total_transacciones": len(ventas),
        "total_articulos": total_cant,
        "total_ventas": total_dia,
        "ventas_resumen": [
            {
                "id": v.id,
                "numero_venta": v.numero_venta,
                "cliente": f"{v.cliente.nombre} {v.cliente.apellido}" if v.cliente else "Cliente General",
                "hora": v.fecha_hora.strftime("%H:%M") if v.fecha_hora else "",
                "items_count": len(v.detalles),
                "total": float(v.total),
                "estado": v.estado,
            }
            for v in ventas
        ],
        "detalles": items_desglosados
    }


@router.get("/diario", summary="Consultar datos del reporte diario de ventas en JSON")
def reporte_diario_json(
    fecha: Optional[str] = Query(None, description="Fecha YYYY-MM-DD (por defecto hoy)"),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role("Administrador", "Empleado")),
):
    fecha_consulta = fecha or date.today().isoformat()
    return _obtener_ventas_dia(fecha_consulta, db)


@router.get("/diario/excel", summary="Exportar reporte diario de ventas en formato Excel (.xlsx)")
def exportar_reporte_diario_excel(
    fecha: Optional[str] = Query(None, description="Fecha YYYY-MM-DD (por defecto hoy)"),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role("Administrador", "Empleado")),
):
    fecha_consulta = fecha or date.today().isoformat()
    datos = _obtener_ventas_dia(fecha_consulta, db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Ventas {fecha_consulta}"
    ws.views.sheetView[0].showGridLines = True

    # Estilos
    title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="3D2314", end_color="3D2314", fill_type="solid")
    sub_font = Font(name="Segoe UI", size=10, italic=True, color="555555")

    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="6F4E37", end_color="6F4E37", fill_type="solid")

    data_font = Font(name="Segoe UI", size=10)
    total_font = Font(name="Segoe UI", size=11, bold=True, color="1F2937")
    total_fill = PatternFill(start_color="EADBC8", end_color="EADBC8", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # Encabezado del reporte
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "CAFÉ SALOMÉ — REPORTE DIARIO DE VENTAS"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:K2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Fecha de Operación: {fecha_consulta} | Generado por: {current_user.nombre} {current_user.apellido} ({current_user.rol.nombre}) | Sistema: Café Cato FullStack"
    sub_cell.font = sub_font
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Resumen rápido
    ws["A4"] = "Transacciones:"
    ws["B4"] = datos["total_transacciones"]
    ws["C4"] = "Artículos Vendidos:"
    ws["D4"] = datos["total_articulos"]
    ws["E4"] = "Total Recaudado:"
    ws["F4"] = f"${datos['total_ventas']:,.2f}"
    for col in ["A", "C", "E"]:
        ws[f"{col}4"].font = Font(name="Segoe UI", size=10, bold=True, color="3D2314")

    # Columnas de tabla
    headers = [
        "Fecha", "Hora", "Nº Venta", "Cliente", "Producto / Servicio",
        "Tipo", "Cant.", "Precio Unit.", "Subtotal", "Método Pago", "Estado"
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [1,2,3,6,7,10,11] else "left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[6].height = 24

    row_idx = 7
    for item in datos["detalles"]:
        ws.cell(row=row_idx, column=1, value=item["fecha"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=item["hora"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3, value=item["numero_venta"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4, value=item["cliente"])
        ws.cell(row=row_idx, column=5, value=item["item"])
        ws.cell(row=row_idx, column=6, value=item["tipo"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=7, value=item["cantidad"]).alignment = Alignment(horizontal="center")
        
        c_price = ws.cell(row=row_idx, column=8, value=item["precio_unitario"])
        c_price.number_format = "$#,##0.00"
        c_price.alignment = Alignment(horizontal="right")

        c_sub = ws.cell(row=row_idx, column=9, value=item["subtotal"])
        c_sub.number_format = "$#,##0.00"
        c_sub.alignment = Alignment(horizontal="right")

        ws.cell(row=row_idx, column=10, value=item["metodo_pago"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=11, value=item["estado"]).alignment = Alignment(horizontal="center")

        for col_idx in range(1, 12):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border

        row_idx += 1

    # Fila de Totales
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
    tot_label = ws.cell(row=row_idx, column=1, value="TOTAL GENERAL DEL DÍA:")
    tot_label.font = total_font
    tot_label.alignment = Alignment(horizontal="right", vertical="center")

    tot_cant = ws.cell(row=row_idx, column=7, value=datos["total_articulos"])
    tot_cant.font = total_font
    tot_cant.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=row_idx, column=8, value="")

    tot_val = ws.cell(row=row_idx, column=9, value=datos["total_ventas"])
    tot_val.font = total_font
    tot_val.number_format = "$#,##0.00"
    tot_val.alignment = Alignment(horizontal="right", vertical="center")

    ws.cell(row=row_idx, column=10, value="")
    ws.cell(row=row_idx, column=11, value="")

    for c in range(1, 12):
        ws.cell(row=row_idx, column=c).fill = total_fill
        ws.cell(row=row_idx, column=c).border = thin_border
    ws.row_dimensions[row_idx].height = 22

    # Ajuste automático de anchos de columna
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.row > 2 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Reporte_Ventas_{fecha_consulta}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/diario/pdf", summary="Exportar reporte diario de ventas en formato PDF estructurado")
def exportar_reporte_diario_pdf(
    fecha: Optional[str] = Query(None, description="Fecha YYYY-MM-DD (por defecto hoy)"),
    db: Session = Depends(get_db),
    current_user: Usuario = Depends(require_role("Administrador", "Empleado")),
):
    fecha_consulta = fecha or date.today().isoformat()
    datos = _obtener_ventas_dia(fecha_consulta, db)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'RepTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor("#3d2314"),
        alignment=1,
        spaceAfter=4
    )
    meta_style = ParagraphStyle(
        'RepMeta',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.HexColor("#6b4c38"),
        alignment=1,
        spaceAfter=12
    )

    elements = []

    # Encabezado
    elements.append(Paragraph("<b>CAFÉ SALOMÉ — REPORTE DIARIO DE VENTAS</b>", title_style))
    elements.append(Paragraph(
        f"Fecha Operativa: <b>{fecha_consulta}</b> | Generado por: {current_user.nombre} {current_user.apellido} ({current_user.rol.nombre}) | Sistema: Café Cato API FullStack",
        meta_style
    ))

    # Tarjetas de resumen superior
    summary_data = [
        [
            f"Transacciones Realizadas: {datos['total_transacciones']}",
            f"Artículos / Servicios Comercializados: {datos['total_articulos']}",
            f"Recaudo Total del Día: ${datos['total_ventas']:,.2f}"
        ]
    ]
    sum_table = Table(summary_data, colWidths=[240, 260, 240])
    sum_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8f3ee")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#c8a78e")),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor("#3d2314")),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(sum_table)
    elements.append(Spacer(1, 14))

    # Tabla de Ventas y Detalles
    headers = ["Hora", "Nº Venta", "Cliente", "Ítem Vendido", "Tipo", "Cant.", "Precio Unit.", "Subtotal", "Método", "Estado"]
    rows = [headers]

    for item in datos["detalles"]:
        rows.append([
            item["hora"],
            item["numero_venta"],
            item["cliente"][:20],
            item["item"][:26],
            item["tipo"],
            str(item["cantidad"]),
            f"${item['precio_unitario']:,.2f}",
            f"${item['subtotal']:,.2f}",
            item["metodo_pago"][:12],
            item["estado"]
        ])

    # Fila total
    rows.append([
        "TOTAL", "", "", "", "",
        str(datos["total_articulos"]),
        "",
        f"${datos['total_ventas']:,.2f}",
        "", ""
    ])

    table = Table(rows, colWidths=[45, 80, 100, 160, 55, 40, 75, 85, 60, 50])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#523321")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 9),
        ('ALIGN', (0,1), (0,-2), 'CENTER'),
        ('ALIGN', (1,1), (1,-2), 'CENTER'),
        ('ALIGN', (5,1), (5,-2), 'CENTER'),
        ('ALIGN', (6,1), (7,-2), 'RIGHT'),
        ('FONTSIZE', (0,1), (-1,-2), 8),
        ('GRID', (0,0), (-1,-2), 0.5, colors.HexColor("#d9cbbe")),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor("#faf7f4")]),
        # Total Row
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor("#e8dacb")),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,-1), (-1,-1), 9),
        ('ALIGN', (5,-1), (5,-1), 'CENTER'),
        ('ALIGN', (7,-1), (7,-1), 'RIGHT'),
        ('GRID', (0,-1), (-1,-1), 1, colors.HexColor("#523321")),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 15))

    footer_style = ParagraphStyle(
        'RepFoot',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.HexColor("#7b614e"),
        alignment=1
    )
    elements.append(Paragraph(f"Documento de control y auditoría contable interna — Café Salomé. Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", footer_style))

    doc.build(elements)
    buffer.seek(0)

    filename = f"Reporte_Ventas_{fecha_consulta}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
